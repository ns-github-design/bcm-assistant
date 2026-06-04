"""Build the AI BCM Assistant 5-minute demo deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PURPLE = RGBColor(0x6B, 0x5B, 0xE3)
PURPLE_LIGHT = RGBColor(0xE6, 0xE2, 0xFA)
INK = RGBColor(0x14, 0x14, 0x1F)
MUTED = RGBColor(0x55, 0x55, 0x66)
BG = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xFF, 0x8A, 0x3D)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x2E, 0x2E)
BORDER = RGBColor(0xE5, 0xE5, 0xEC)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    return s


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_orb(slide, cx, cy, r, color=PURPLE_LIGHT):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    o.line.fill.background()
    o.fill.solid()
    o.fill.fore_color.rgb = color
    return o


def add_rect(slide, x, y, w, h, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    r.line.fill.background()
    r.fill.solid()
    r.fill.fore_color.rgb = color
    return r


def add_card(slide, x, y, w, h, title, body, accent=PURPLE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.adjustments[0] = 0.06
    card.line.color.rgb = BORDER
    card.line.width = Pt(0.75)
    card.fill.solid()
    card.fill.fore_color.rgb = BG
    add_rect(slide, x, y, Inches(0.08), h, accent)
    add_text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.4),
             Inches(0.5), title, size=15, bold=True, color=INK)
    add_text(slide, x + Inches(0.25), y + Inches(0.65), w - Inches(0.4),
             h - Inches(0.7), body, size=11, color=MUTED)


def slide_footer(slide, num, total):
    add_text(slide, Inches(0.5), Inches(7.05), Inches(8), Inches(0.3),
             "AI BCM ASSISTANT  ·  Devi Priya & Nasreen Sarah",
             size=9, color=MUTED)
    add_text(slide, Inches(11.8), Inches(7.05), Inches(1.2), Inches(0.3),
             f"{num} / {total}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)


# ---------- Slide 1: Title ----------
s = add_slide()
add_orb(s, Inches(4.0), Inches(3.5), Inches(2.6), RGBColor(0xEC, 0xE7, 0xFD))
add_orb(s, Inches(3.7), Inches(3.4), Inches(1.9), RGBColor(0xC7, 0xBE, 0xF8))
add_orb(s, Inches(3.5), Inches(3.3), Inches(1.1), RGBColor(0x8E, 0x7B, 0xE5))
add_text(s, Inches(1.0), Inches(2.1), Inches(8), Inches(0.4),
         "Presented by Devi Priya & Nasreen Sarah", size=16, color=INK)
add_rect(s, Inches(1.0), Inches(2.9), Emu(15875), Inches(2.6), INK)
add_text(s, Inches(1.4), Inches(3.05), Inches(11), Inches(2),
         "AI BCM ASSISTANT", size=66, bold=True, color=INK, font="Cambria")
add_text(s, Inches(1.4), Inches(5.3), Inches(11), Inches(0.5),
         "An agentic copilot for Business Continuity & FinOps on Azure",
         size=18, color=MUTED)

# ---------- Slide 2: Problem ----------
s = add_slide()
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "THE PROBLEM", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1),
         "DR planning is slow, manual, and expensive.", size=32, bold=True)
add_text(s, Inches(0.7), Inches(1.85), Inches(12), Inches(0.6),
         "Banks pay for Active-Active DR they don't always need — and still can't prove BCM compliance.",
         size=15, color=MUTED)

cards = [
    ("Spreadsheets & silos",
     "Architects, FinOps and BCM teams work in separate tools. Tier, RTO, RPO and "
     "cost decisions live in disconnected SMDB / CMDB / budget sheets."),
    ("Standby idle 99% of the time",
     "DR environments are provisioned identical to prod and run 24×7 — burning "
     "budget on capacity that is only used during a real failover."),
    ("Compliance as an afterthought",
     "Pilot Light, Warm Standby, Active-Active — choosing the right pattern per "
     "service is a multi-week consulting exercise. Audits suffer."),
    ("No single source of truth",
     "Cost Mgmt, ARM, Entra, BCM Catalogue, Approval Portal — five tools, "
     "five exports, zero traceability."),
]
for i, (t, b) in enumerate(cards):
    col, row = i % 2, i // 2
    add_card(s, Inches(0.7 + col * 6.1), Inches(2.7 + row * 2),
             Inches(5.9), Inches(1.8), t, b)
slide_footer(s, 2, 9)

# ---------- Slide 3: Solution ----------
s = add_slide()
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "THE SOLUTION", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1),
         "An agent that walks you from inventory to approved DR plan.",
         size=28, bold=True)
add_text(s, Inches(0.7), Inches(1.95), Inches(12), Inches(0.6),
         "Conversational. Auditable. Backed by your real Azure data and BCM policy.",
         size=15, color=MUTED)
vp = [
    ("Conversational",
     "Plain-English chat replaces six dashboards. The agent asks for the service "
     "name and subscription — the rest is automatic."),
    ("Pilot Light by default",
     "Per-resource DR strategy: pre-provision, replicate, rebuild on demand or "
     "native DR — costed and RTO-checked."),
    ("Stage-driven",
     "Six guided stages: Analyse, Set RTO/RPO, Adopt Pilot Light, Cost & Approval, "
     "BCM Test, Go to Production."),
]
for i, (t, b) in enumerate(vp):
    add_card(s, Inches(0.7 + i * 4.1), Inches(2.9), Inches(3.95), Inches(3.6),
             t, b, accent=ACCENT if i == 1 else PURPLE)
slide_footer(s, 3, 9)

# ---------- Slide 4: Architecture ----------
s = add_slide()
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "ARCHITECTURE", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1),
         "Three tiers, one agent loop.", size=30, bold=True)


def arch_box(slide, x, y, w, h, title, sub, accent=PURPLE):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.adjustments[0] = 0.08
    box.line.color.rgb = accent
    box.line.width = Pt(1.25)
    box.fill.solid()
    box.fill.fore_color.rgb = BG
    add_text(slide, x + Inches(0.2), y + Inches(0.18), w - Inches(0.4),
             Inches(0.4), title, size=14, bold=True, color=INK)
    add_text(slide, x + Inches(0.2), y + Inches(0.62), w - Inches(0.4),
             h - Inches(0.7), sub, size=10.5, color=MUTED)


arch_box(s, Inches(0.7), Inches(2.2), Inches(3.7), Inches(2.4),
         "Frontend  ·  React + Vite",
         "• Chat UI with 6 workflow stages\n"
         "• Live SSE stream of agent events\n"
         "• Right-side service context panel\n"
         "• Demo mode for offline walkthroughs",
         accent=PURPLE)
arch_box(s, Inches(4.85), Inches(2.2), Inches(3.7), Inches(2.4),
         "Backend  ·  FastAPI",
         "• /api/assess — SSE assessment stream\n"
         "• /api/chat — follow-up Q&A\n"
         "• Pilot Light strategy engine\n"
         "• Pydantic schemas, mock + live modes",
         accent=ACCENT)
arch_box(s, Inches(9.0), Inches(2.2), Inches(3.7), Inches(2.4),
         "Agent  ·  GitHub Copilot SDK",
         "• CopilotClient over JSON-RPC\n"
         "• gpt-4.1 with streaming reasoning\n"
         "• 4 typed tools (Pydantic)\n"
         "• System prompt = BCM playbook",
         accent=PURPLE)

# arrows
for x1, x2 in [(Inches(4.4), Inches(4.85)), (Inches(8.55), Inches(9.0))]:
    line = s.shapes.add_connector(2, x1, Inches(3.4), x2, Inches(3.4))
    line.line.color.rgb = MUTED
    line.line.width = Pt(1.5)

arch_box(s, Inches(0.7), Inches(4.95), Inches(12.0), Inches(1.85),
         "Data sources & tools",
         "Tools:  assess_azure_resources  ·  get_budget_status  ·  query_cmdb  ·  suggest_pilot_light_plan\n"
         "Sources:  Azure ARM  ·  Cost Mgmt  ·  Entra ID  ·  SMDB  ·  CMDB  ·  BCM Catalogue  ·  Approval Portal\n"
         "Today: mock-first demo data    →    Tomorrow: live Azure SDK + MCP server",
         accent=PURPLE)
slide_footer(s, 4, 9)

# ---------- Slide 5: Stages ----------
s = add_slide()
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "THE WORKFLOW", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1),
         "Six guided stages, one conversation.", size=30, bold=True)
stages = [
    ("1", "Analyse my service",
     "Connect to Azure, scan subs & assets, full financial comparison."),
    ("2", "Set RTO & RPO",
     "Pull SMDB, validate tier, set RPO in business language."),
    ("3", "Adopt Pilot Light",
     "Fit-gap vs BCM Catalogue, deliverables per service."),
    ("4", "Cost & approval",
     "Final cost, savings, stakeholder sign-off."),
    ("5", "BCM test",
     "Auto test plan, failover execution, compliance report."),
    ("6", "Go to production",
     "Pre-prod checklist, go-live, rollback plan."),
]
for i, (n, t, d) in enumerate(stages):
    col, row = i % 3, i // 3
    x, y = Inches(0.7 + col * 4.1), Inches(2.4 + row * 2.0)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y,
                             Inches(3.95), Inches(1.75))
    box.adjustments[0] = 0.08
    box.line.color.rgb = BORDER
    box.line.width = Pt(0.75)
    box.fill.solid()
    box.fill.fore_color.rgb = BG
    circle = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.2),
                                y + Inches(0.22), Inches(0.55), Inches(0.55))
    circle.line.fill.background()
    circle.fill.solid()
    circle.fill.fore_color.rgb = PURPLE if i < 3 else PURPLE_LIGHT
    add_text(s, x + Inches(0.2), y + Inches(0.22), Inches(0.55), Inches(0.55),
             n, size=18, bold=True,
             color=BG if i < 3 else PURPLE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.95), y + Inches(0.25), Inches(2.9), Inches(0.5),
             t, size=14, bold=True)
    add_text(s, x + Inches(0.25), y + Inches(0.95), Inches(3.55), Inches(0.8),
             d, size=10.5, color=MUTED)
slide_footer(s, 5, 9)

# ---------- Slide 6: Live Demo ----------
s = add_slide()
add_orb(s, Inches(10.5), Inches(3.5), Inches(2.4), RGBColor(0xEC, 0xE7, 0xFD))
add_orb(s, Inches(10.5), Inches(3.5), Inches(1.6), RGBColor(0xC7, 0xBE, 0xF8))
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "LIVE DEMO", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(1.2), Inches(11), Inches(1.4),
         "Live demo", size=66, bold=True, font="Cambria")
add_text(s, Inches(0.7), Inches(2.6), Inches(10), Inches(0.5),
         "IB Investment Research API  ·  subscription UBS-IB-PROD-WEU",
         size=16, color=MUTED)
demo = [
    ("1", "Analyse the service", "Subscriptions, assets and a full financial picture in seconds."),
    ("2", "See the pain",        "DR standby running 24/7, $100k over budget."),
    ("3", "Adopt Pilot Light",   "Per-service strategy + $545k projected annual saving."),
    ("4", "Agent reasoning",     "Watch the thinking trace and the right-side context update live."),
]
y = Inches(3.4)
for n, t, d in demo:
    add_text(s, Inches(0.7), y, Inches(0.5), Inches(0.5),
             n, size=22, bold=True, color=PURPLE)
    add_text(s, Inches(1.2), y + Inches(0.05), Inches(3.2), Inches(0.4),
             t, size=14, bold=True)
    add_text(s, Inches(4.4), y + Inches(0.08), Inches(7), Inches(0.5),
             d, size=12, color=MUTED)
    y += Inches(0.7)
slide_footer(s, 6, 9)

# ---------- Slide 7: Demo highlights ----------
s = add_slide()
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "DEMO HIGHLIGHTS", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1),
         "What you'll see on screen.", size=30, bold=True)


def kpi(slide, x, y, w, h, value, label, color=PURPLE):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.adjustments[0] = 0.08
    box.line.color.rgb = BORDER
    box.line.width = Pt(0.75)
    box.fill.solid()
    box.fill.fore_color.rgb = BG
    add_text(slide, x, y + Inches(0.35), w, Inches(0.9),
             value, size=34, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + Inches(1.3), w, Inches(0.5),
             label, size=12, color=MUTED, align=PP_ALIGN.CENTER)


kpi(s, Inches(0.7),  Inches(2.4), Inches(3.0), Inches(2.0), "$1.3M",  "Current annual spend", color=INK)
kpi(s, Inches(3.85), Inches(2.4), Inches(3.0), Inches(2.0), "108%",   "Budget utilisation", color=RED)
kpi(s, Inches(7.0),  Inches(2.4), Inches(3.0), Inches(2.0), "$755k",  "Pilot Light target / yr", color=PURPLE)
kpi(s, Inches(10.15),Inches(2.4), Inches(2.5), Inches(2.0), "↓ $545k","Projected annual saving", color=GREEN)

add_text(s, Inches(0.7), Inches(4.7), Inches(12), Inches(0.4),
         "Key findings the agent surfaces", size=14, bold=True)
findings = [
    ("!", "DR standby (NEU) runs 24/7 — idle 99% of the time, costs the same as prod.", RED),
    ("!", "$100k over budget at 108% of approved $1.2M.", RED),
    ("✓", "Key Vault + Entra ID — always-on, low cost, no action needed.", GREEN),
    ("✓", "Storage GRS — geo-redundancy built-in, fully Pilot Light compatible.", GREEN),
]
y = Inches(5.15)
for icon, txt, col in findings:
    add_text(s, Inches(0.7), y, Inches(0.4), Inches(0.4), icon,
             size=14, bold=True, color=col)
    add_text(s, Inches(1.1), y + Inches(0.02), Inches(11.5), Inches(0.4),
             txt, size=12)
    y += Inches(0.4)
slide_footer(s, 7, 9)

# ---------- Slide 8: Tech & Roadmap ----------
s = add_slide()
add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4),
         "TECH & ROADMAP", size=11, bold=True, color=PURPLE)
add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1),
         "Built on the GitHub Copilot SDK.", size=30, bold=True)

add_text(s, Inches(0.7), Inches(2.0), Inches(6), Inches(0.4),
         "Tech stack", size=14, bold=True, color=PURPLE)
techs = [
    "Frontend   ·  React 18, Vite, Zustand, Tailwind, react-router",
    "Backend    ·  FastAPI, Pydantic v2, sse-starlette, Python 3.11",
    "Agent      ·  GitHub Copilot SDK + CLI (JSON-RPC), gpt-4.1",
    "Tools      ·  4 typed tools with progress streaming",
    "Hosting    ·  Azure Container Apps (frontend + backend)",
    "Future     ·  Live Azure SDK discovery, MCP server",
]
y = Inches(2.5)
for t in techs:
    add_text(s, Inches(0.7), y, Inches(6), Inches(0.35), "•  " + t, size=12)
    y += Inches(0.4)

add_text(s, Inches(7.2), Inches(2.0), Inches(6), Inches(0.4),
         "What's next", size=14, bold=True, color=ACCENT)
roadmap = [
    ("Live Azure discovery",
     "Replace mock data with ARM + Cost Mgmt SDK calls."),
    ("Stages 4–6",
     "Cost & approval, BCM test execution, go-live checklist."),
    ("MCP server",
     "Expose tools to other agents via Model Context Protocol."),
    ("Persistent sessions",
     "Multi-turn memory and audit trail per service."),
]
y = Inches(2.5)
for t, d in roadmap:
    add_text(s, Inches(7.2), y, Inches(5.5), Inches(0.35),
             "›  " + t, size=12, bold=True)
    add_text(s, Inches(7.45), y + Inches(0.32), Inches(5.3), Inches(0.4),
             d, size=10.5, color=MUTED)
    y += Inches(0.85)
slide_footer(s, 8, 9)

# ---------- Slide 9: Thank you ----------
s = add_slide()
add_orb(s, Inches(6.66), Inches(3.6), Inches(2.8), RGBColor(0xEC, 0xE7, 0xFD))
add_orb(s, Inches(6.66), Inches(3.6), Inches(1.9), RGBColor(0xC7, 0xBE, 0xF8))
add_orb(s, Inches(6.66), Inches(3.6), Inches(1.0), RGBColor(0x8E, 0x7B, 0xE5))
add_text(s, Inches(0), Inches(2.4), Inches(13.33), Inches(1.6),
         "Thank you.", size=72, bold=True, color=INK,
         align=PP_ALIGN.CENTER, font="Cambria")
add_text(s, Inches(0), Inches(4.2), Inches(13.33), Inches(0.5),
         "Questions?", size=22, color=MUTED, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(6.4), Inches(13.33), Inches(0.4),
         "Devi Priya  ·  Nasreen Sarah", size=13, color=MUTED,
         align=PP_ALIGN.CENTER)

out = r"c:\Users\nasreensarah\bcm-assistant\docs\AI_BCM_Assistant_Demo.pptx"
prs.save(out)
print(f"Saved: {out}")
